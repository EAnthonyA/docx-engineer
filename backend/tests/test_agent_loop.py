from pathlib import Path
import asyncio

import pytest
from docx import Document
from docx.shared import Inches
from docxkit import DocxTools
from fastapi import BackgroundTasks, HTTPException

from app import jobs, main


@pytest.fixture
def job(tmp_path, monkeypatch):
    monkeypatch.setattr(jobs, "JOBS_DIR", tmp_path)
    monkeypatch.setattr(main, "JOBS_DIR", tmp_path)
    monkeypatch.setattr(jobs, "_jobs", {})
    job = jobs.create_job("Set font size to 24 pt")
    doc = Document()
    doc.add_paragraph("Hello world")
    doc.save(job.input_path)
    return job


@pytest.mark.parametrize("change", ["size", "layout"])
def test_successful_edits_finish_on_first_attempt(job, monkeypatch, change):
    calls = []
    def generate(*args):
        calls.append(1)
        return "def edit(doc, tools):\n    pass"
    monkeypatch.setattr(main.llm, "generate_script", generate)
    def sandbox(job_id):
        doc = Document(job.input_path)
        if change == "size":
            DocxTools().set_format(doc.paragraphs[0], size_pt=24)
        else:
            doc.sections[0].left_margin = Inches(2)
        doc.save(Path(job.input_path).parent / "out/out.docx")
        return {"success": True}
    monkeypatch.setattr(main.sandbox, "run_script", sandbox)
    main._run_agent_loop(job.id, clarify=False)
    assert job.status == "needs_review"
    assert job.attempt == 1
    assert len(calls) == 1
    assert job.diff["package_changed"] is True


def test_auto_mode_runs_a_valid_plan_in_rust(job, monkeypatch):
    monkeypatch.setenv("DOCX_ENGINE_MODE", "auto")
    monkeypatch.setattr(main.llm, "generate_plan", lambda *args: """{
      \"schema_version\": 1,
      \"operations\": [{
        \"id\": \"format_b_tags\", \"type\": \"format_b_tags\", \"locations\": [\"body\"],
        \"open_tag\": \"<B>\", \"close_tag\": \"<D>\", \"bold\": true
      }]
    }""")
    monkeypatch.setattr(main.llm, "generate_script", lambda *args: pytest.fail("Python generation should not run"))

    def sandbox(job_id, *, engine):
        assert engine == "rust"
        assert (Path(job.input_path).parent / "plan.json").exists()
        doc = Document(job.input_path)
        doc.paragraphs[0].add_run(" changed")
        doc.save(Path(job.input_path).parent / "out/out.docx")
        return {"success": True, "engine": "rust"}

    monkeypatch.setattr(main.sandbox, "run_script", sandbox)
    main._run_agent_loop(job.id, clarify=False)
    assert job.status == "needs_review"
    assert any(entry["engine"] == "rust" for entry in job.diagnostics if entry["event"] == "execution_prepared")


def test_auto_mode_falls_back_once_from_rust_to_python(job, monkeypatch):
    monkeypatch.setenv("DOCX_ENGINE_MODE", "auto")
    monkeypatch.setattr(main.llm, "generate_plan", lambda *args: """{
      \"schema_version\": 1, \"operations\": [{
        \"id\": \"format_b_tags\", \"type\": \"format_b_tags\", \"locations\": [\"body\"],
        \"open_tag\": \"<B>\", \"close_tag\": \"<D>\", \"bold\": true
      }]
    }""")
    monkeypatch.setattr(main.llm, "generate_script", lambda *args: "def edit(doc, tools):\n    pass")
    calls = []

    def sandbox(job_id, *, engine="python"):
        calls.append(engine)
        if engine == "rust":
            return {"success": False, "error": "unsupported_content", "error_kind": "rust_error"}
        doc = Document(job.input_path)
        doc.paragraphs[0].add_run(" changed")
        doc.save(Path(job.input_path).parent / "out/out.docx")
        return {"success": True}

    monkeypatch.setattr(main.sandbox, "run_script", sandbox)
    main._run_agent_loop(job.id, clarify=False)
    assert calls == ["rust", "python"]
    assert job.status == "needs_review"


@pytest.mark.parametrize("kind,stage,expected_attempts", [
    ("executor_error", "starting_container", 1),
    ("executor_timeout", "unknown", 1),
    ("document_error", "loading_document", 1),
    ("oom", "loading_document", 1),
    ("oom", "editing_document", 2),
    ("timeout", "saving_document", 1),
    ("timeout", "editing_document", 2),
    ("script_error", "editing_document", 5),
])
def test_retry_policy_depends_on_failure_kind(job, monkeypatch, kind, stage, expected_attempts):
    calls = []
    def generate(*args):
        calls.append(1)
        return "def edit(doc, tools):\n    pass"
    monkeypatch.setattr(main.llm, "generate_script", generate)
    monkeypatch.setattr(main.sandbox, "run_script", lambda _: {
        "success": False, "error": "diagnostic", "error_kind": kind, "stage": stage,
    })
    main._run_agent_loop(job.id, clarify=False)
    assert job.status == "stuck"
    assert len(calls) == expected_attempts
    assert job.attempt == expected_attempts


def test_genuine_no_op_is_not_accepted_as_a_success(job, monkeypatch):
    monkeypatch.setattr(main.llm, "generate_script", lambda *args: "def edit(doc, tools):\n    pass")
    def sandbox(_):
        doc = Document(job.input_path)
        doc.save(Path(job.input_path).parent / "out/out.docx")
        return {"success": True}
    monkeypatch.setattr(main.sandbox, "run_script", sandbox)
    main._run_agent_loop(job.id, clarify=False)
    assert job.status == "stuck"
    assert job.attempt == 5
    assert "no changes" in job.attempt_error


def test_refinement_is_sent_to_model_even_when_no_script_was_generated(job):
    job.status = "stuck"
    asyncio.run(main.refine_job(job.id, main.RefineRequest(note="Use 12 pt instead"), BackgroundTasks(), True))
    assert ("Additional editing instruction", "Use 12 pt instead") in job.clarifications


def test_archived_jobs_cannot_schedule_a_run_without_a_document(job):
    job.status = "stuck"
    job.input_path = ""
    with pytest.raises(HTTPException) as error:
        asyncio.run(main.refine_job(job.id, main.RefineRequest(note="Try again"), BackgroundTasks(), True))
    assert error.value.status_code == 409
