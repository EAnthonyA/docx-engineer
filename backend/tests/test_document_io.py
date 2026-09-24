import io
import struct
import zipfile
import zlib

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE
from lxml import etree

from document_io import save_document


def test_streaming_writer_preserves_every_package_part(tmp_path):
    doc = Document()
    para = doc.add_paragraph("before\tand after\nnext line")
    para.runs[0].bold = True
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "table"
    doc.sections[0].header.paragraphs[0].text = "header"
    def chunk(kind, data):
        return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data))
    # One RGB pixel; exercise a real binary image part and its relationships.
    image = (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0))
             + chunk(b"IDAT", zlib.compress(b"\x00\xff\x00\x00")) + chunk(b"IEND", b""))
    doc.add_picture(io.BytesIO(image))
    relationship = doc.part.relate_to("https://example.com", RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), relationship)
    para._p.append(link)
    normal, streamed = tmp_path / "normal.docx", tmp_path / "streamed.docx"
    doc.save(normal)
    save_document(doc, streamed)
    with zipfile.ZipFile(normal) as before, zipfile.ZipFile(streamed) as after:
        assert set(before.namelist()) == set(after.namelist())
        for name in before.namelist():
            a, b = before.read(name), after.read(name)
            if name.endswith((".xml", ".rels")):
                assert etree.tostring(etree.fromstring(a), method="c14n") == etree.tostring(etree.fromstring(b), method="c14n")
            else:
                assert a == b
    result = Document(streamed)
    assert result.paragraphs[0].text == para.text
    assert result.sections[0].header.paragraphs[0].text == "header"
    assert len(result.inline_shapes) == 1
    from app.docx_inspect import compute_diff
    assert compute_diff(str(normal), str(streamed))["package_changed"] is False
